"""Markdown → .docx (python-docx) with a small, dependable subset: headings, bullets, numbered
lists, bold/italic/code inline, pipe tables, block quotes."""

from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def _add_inline(paragraph, text: str) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _flush_table(doc, rows: list[list[str]]) -> None:
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in r)]
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, r in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            _add_inline(p, r[j].strip() if j < len(r) else "")
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def markdown_to_docx(markdown: str, *, title: str | None = None) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    if title:
        h = doc.add_heading(title, level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_rows: list[list[str]] = []
    in_code = False
    code_buf: list[str] = []
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                r = p.add_run("\n".join(code_buf))
                r.font.name = "Consolas"
                r.font.size = Pt(9)
                code_buf = []
            in_code = not in_code
            continue
        if in_code:
            code_buf.append(raw)
            continue
        if line.strip().startswith("|"):
            table_rows.append([c for c in line.strip().strip("|").split("|")])
            continue
        if table_rows:
            _flush_table(doc, table_rows)
            table_rows = []
        if not line.strip():
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = min(len(m.group(1)), 4)
            doc.add_heading(re.sub(r"[*`]", "", m.group(2)), level=level)
            continue
        if re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, re.sub(r"^\s*[-*]\s+", "", line))
            continue
        if re.match(r"^\s*\d+[.)]\s+", line):
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, re.sub(r"^\s*\d+[.)]\s+", "", line))
            continue
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _add_inline(p, line.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
            continue
        p = doc.add_paragraph()
        _add_inline(p, line)
    if table_rows:
        _flush_table(doc, table_rows)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
